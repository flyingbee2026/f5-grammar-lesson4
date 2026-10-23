#!/usr/bin/env python3
"""Build F5 Grammar Lesson 4 (revised v4): Cleft sentences — teacher version.

Template = Helen's finalised Lesson 3 TEACHER docx (keeps styles.xml).
"""
import copy
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

MODEL = "lesson3-teacher.docx"
OUT = {"teacher": "Lesson 4 - Cleft sentences (Teacher's).docx"}
XMLSPACE = "{http://www.w3.org/XML/1998/namespace}space"

# ---------------------------------------------------------------- model indices
H_SCHOOL, H_SUBJ, H_TITLE, H_TOPIC = 0, 1, 2, 3
P_H2 = 4            # Heading 2 part heading
P_PLAIN = 5         # plain Normal body paragraph
T_QUIZ = 6          # 5x1 numbered quiz table
P_BLANK = 7         # empty Normal paragraph
P_ANS_LABEL = 8     # bold 'Teacher's answers'
T_ANSWERS = 9       # 6x2 F2F2F2 answers table
P_H2B = 11
P_INSTR = 12        # plain instruction line
T_WORDS = 13        # 4-col warm-up table
P_H3 = 17           # Heading 3 example heading (single run)
T_PAIR = 18         # 2x2 DEEBF7 Plain / More formal table
P_BOX = 19          # bordered demo paragraph
P_H3_LINK = 30      # 'Nominalisation as a linking device'
P_PLAIN2 = 31
T_LINK = 32         # 2x2 DEEBF7 linking table
P_H1 = 35           # Heading 1 'Part 3: Practice sets'
P_SETINTRO = 36
P_SETTITLE = 37     # Heading 3 set title (single run)
T_SET = 38          # 3x2 set table
P_H1W = 53          # Heading 1 'Part 4: Writing Task'
P_INSTR_B = 54      # lavender bold 'Instructions:'
P_INSTR_1 = 55      # lavender task line
P_INSTR_2 = 56      # lavender bullet line (wingdings sym + br)
P_SAMPLE = 58       # grey sample box: bold label run + body run
P_NOTE = 61         # grey teacher note paragraph

# ---------------------------------------------------------------- header
SCHOOL = "St. Joseph\u2019s Anglo-Chinese School"
SUBJ = "F.5 English Language"
TITLE = "Grammar \u2013 Lesson 4: Cleft sentences"
TOPIC = "Cleft sentences"

# ---------------------------------------------------------------- Part 1: Proofreading
QUIZ_INTRO = ("Your teacher has put two paragraphs on the board. One argues that "
              "technology has made people lazier (Paragraph A). The other argues that "
              "technology has not made us lazier (Paragraph B). Each paragraph contains "
              "FIVE errors. The errors include relative clauses, inverted patterns and "
              "nominalisation. Proofread the paragraphs.")

PARA_A = (
    "Everyone claims that technology makes life easier, but the truth is far less pleasant. "
    "Not only teenagers order food instead of cooking, but they also replace real "
    "conversations with emoji-filled messages. They sit for hours on apps, that were designed "
    "to save time, and end up achieving nothing meaningful. The convenient of instant "
    "delivery has destroyed any reason to walk to the shops. Rarely parents notice that their "
    "children have forgotten how to hold a proper conversation. Only if teenagers put down their "
    "phones the family can have a real conversation."
)

PARA_B = (
    "Has technology really destroyed our generation, or have adults simply misunderstood "
    "how we work? What they refuse to accept is that young people today manage schoolwork "
    "alongside digital projects what they could never have completed alone. The believe that "
    "real work must involve physical sweat ignores the fact that running a website demands "
    "just as much effort. Rarely critics acknowledge that a teenager editing videos is working "
    "just as hard as someone lifting boxes. The reason because teenagers seem distracted is not "
    "laziness but the creative work they are doing on their phones. Only when adults look at "
    "the results of their work they can understand what teenagers actually achieve."
)

QUIZ_KEY = [
    ("1A", "Not only do teenagers order food instead of cooking... "
           "(Move 'do' before the subject \u2014 'Not only' inverts the subject and the verb.)"),
    ("2A", "They sit for hours on apps, which were designed to save time... "
           "('that' cannot be used after a comma in a non-restrictive clause; use 'which'.)"),
    ("3A", "The convenience of instant delivery has destroyed any reason to walk to the shops. "
           "('convenient' \u2192 'convenience' \u2014 a noun is needed after 'the'.)"),
    ("4A", "Rarely do parents notice that their children have forgotten how to hold a proper conversation. "
           "(Move 'do' before the subject \u2014 'Rarely' inverts the subject and the verb.)"),
    ("5A", "Only if teenagers put down their phones can the family have a real conversation. "
           "(Move 'can' before the subject \u2014 'Only if' inverts the subject and the verb.)"),
    ("1B", "...digital projects which they could never have completed alone... "
           "('what' cannot be used as a relative pronoun after a noun; use 'which' or 'that'.)"),
    ("2B", "The belief that real work must involve physical sweat ignores the fact that running a website "
           "demands just as much effort. "
           "('believe' \u2192 'belief' \u2014 a noun is needed after 'the'.)"),
    ("3B", "Rarely do critics acknowledge that a teenager editing videos is working just as hard as "
           "someone lifting boxes. "
           "(Move 'do' before the subject \u2014 'Rarely' inverts the subject and the verb.)"),
    ("4B", "The reason why teenagers seem distracted is not laziness but the creative work they are doing "
           "on their phones. "
           "('because' cannot follow 'The reason'; use 'why' or 'that'.)"),
    ("5B", "Only when adults look at the results of their work can they understand what teenagers actually "
           "achieve. "
           "(Move 'can' before the subject \u2014 'Only when' inverts the subject and the verb.)"),
]

# ---------------------------------------------------------------- Part 2: Cleft patterns + class practice
# Order: Wh-clefts first, then It-clefts
P2_INTRO = ("A cleft sentence breaks one idea into two parts so the most important "
            "information stands out. Study the patterns below.")

EXAMPLES = [
    # WH-CLEFTS FIRST
    ("Example 1: Wh-cleft \u2013 presenting a recommendation",
     "Schools should teach students how to communicate face to face.",
     "What schools should do is teach students how to communicate face to face.",
     "Many teachers complain that their students cannot hold a proper conversation. What "
     "schools should do is teach students how to communicate face to face. A single "
     "workshop on active listening would help teenagers rebuild the skills that messaging "
     "apps have quietly taken away."),
    ("Example 2: All-cleft \u2013 naming the only thing needed",
     "Teenagers only need one real conversation to remember the value of talking.",
     "All teenagers need is one real conversation to remember the value of talking.",
     "Parents often buy expensive devices or hire tutors to help their children socially. All "
     "teenagers need is one real conversation to remember the value of talking. Once they "
     "experience the warmth of genuine human connection, the cold comfort of emojis will "
     "seem far less appealing."),
    # THEN IT-CLEFTS
    ("Example 3: It-cleft \u2013 focusing on the cause",
     "People communicate less because they use messaging apps.",
     "It is messaging apps that make people communicate less.",
     "Young people used to meet at the park after school, but now everyone sends messages "
     "from their bedrooms. A quick text is easier than a real conversation, so face-to-face "
     "friendships have become rare. It is messaging apps that make people communicate less, "
     "and this change will only grow deeper as technology advances."),
    ("Example 4: It-cleft \u2013 setting up a contrast",
     "Some people blame social media, but smartphones cause the real problem.",
     "While some people blame social media, it is smartphones that cause the real problem.",
     "Every day there are news stories about the dangers of social media. While some people "
     "blame social media, it is smartphones that cause the real problem. Without a phone in "
     "their pocket, teenagers could not check their accounts every five minutes."),
]

# Class practice after each example (easy, harder)
CLASS_PRACTICE = [
    # After Example 1 (Wh-cleft recommendation)
    [
        ("Class question 1 (Easy): Wh-cleft \u2013 recommendation",
         "Schools should teach students how to start a real conversation.",
         "Begin with \u2018What schools should do \u2026\u2019.",
         "What schools should do is teach students how to start a real conversation."),
        ("Class question 2 (Harder): Wh-cleft \u2013 recommendation",
         "Teenagers need a safe space to practise speaking in groups.",
         "Begin with \u2018What teenagers need \u2026\u2019.",
         "What teenagers need is a safe space to practise speaking in groups."),
    ],
    # After Example 2 (All-cleft)
    [
        ("Class question 1 (Easy): All-cleft",
         "Teenagers only need one supportive friend to feel understood.",
         "Begin with \u2018All teenagers need \u2026\u2019.",
         "All teenagers need is one supportive friend to feel understood."),
        ("Class question 2 (Harder): All-cleft",
         "Teenagers only need to feel brave for a moment to speak to a stranger.",
         "Begin with \u2018All teenagers need \u2026\u2019.",
         "All teenagers need is the bravery to speak to a stranger."),
    ],
    # After Example 3 (It-cleft cause)
    [
        ("Class question 1 (Easy): It-cleft \u2013 the cause",
         "Young people forget social skills because they use messaging apps.",
         "Begin with \u2018It is messaging apps \u2026\u2019.",
         "It is messaging apps that make young people forget social skills."),
        ("Class question 2 (Harder): It-cleft \u2013 the cause",
         "Teenagers struggle to focus because notifications interrupt them constantly.",
         "Begin with \u2018It is the constant interruptions \u2026\u2019.",
         "It is the constant interruptions from notifications that make teenagers struggle to focus."),
    ],
    # After Example 4 (It-cleft contrast)
    [
        ("Class question 1 (Easy): It-cleft \u2013 the contrast",
         "Some blame phones for tiredness, but late-night gaming is the real cause.",
         "Begin with \u2018While some blame phones for tiredness \u2026\u2019.",
         "While some blame phones for tiredness, it is late-night gaming that is the real cause."),
        ("Class question 2 (Harder): It-cleft \u2013 the contrast",
         "People say teenagers are antisocial, but they themselves avoid family dinners.",
         "Begin with \u2018While people say teenagers are antisocial \u2026\u2019.",
         "While people say teenagers are antisocial, it is adults who avoid family dinners."),
    ],
]

# ---------------------------------------------------------------- Part 3: Independent practice sets (4 only, no hints)
SETS_INTRO = ("Rewrite each sentence as a cleft sentence. The type of cleft to use is shown "
              "in brackets.")

SETS = [
    ("Set 1: It-cleft \u2013 the cause",
     "Teenagers send short messages because they are afraid of real conversations.",
     "Use an IT-cleft.",
     "It is the fear of real conversations that makes teenagers send short messages."),
    ("Set 2: Wh-cleft \u2013 the recommendation",
     "Students need to practise face-to-face communication.",
     "Use a WHAT-cleft.",
     "What students need to do is practise face-to-face communication."),
    ("Set 3: All-cleft \u2013 the only thing needed",
     "Teenagers only need to be confident to start a real conversation.",
     "Use an ALL-cleft.",
     "All teenagers need is the confidence to start a real conversation."),
    ("Set 4: It-cleft \u2013 the contrast",
     "People think technology is convenient, but it actually destroys our patience.",
     "Use an IT-cleft with WHILE.",
     "While people think technology is convenient, it is the convenience of modern devices that destroys our patience."),
]

# ---------------------------------------------------------------- Part 4: Writing Task
TASK_TITLE = "Part 4: Writing Task"
TASK_LINE = ("1. Write ONE paragraph (about 80 words) explaining the influence of artificial "
             "intelligence on students, and suggesting ONE thing schools could do to help.")
TASK_BULLETS = [
    ("\u2022", None),
    (" Use ", "plain"),
    ("ONE cleft sentence", "bold_u"),
    (" (It is \u2026 that \u2026 / What \u2026 is \u2026 / All \u2026 needs to do is \u2026).",
     "plain"),
    (None, "br"),
    ("\u2022", None),
    (" Use a relative clause or an inverted pattern e.g. Not only \u2026 / Never \u2026 / "
     "Rarely \u2026", "plain"),
]

SAMPLES = [
    ("Sample answer 1 (It-cleft) ",
     "Artificial intelligence is changing the way students learn, but not every change is "
     "positive. It is the loss of independent thinking that worries teachers most, as too "
     "many teenagers now let apps write their essays without even reading the questions "
     "first. Schools should therefore limit AI use in all homework tasks, which would force "
     "students to develop their own ideas and opinions. A simple rule requiring a first "
     "draft without AI would restore the value of real effort."),
    ("Sample answer 2 (Wh-cleft) ",
     "What concerns parents most is that their children are losing the ability to think for "
     "themselves. Artificial intelligence can answer almost any question instantly, so "
     "students no longer wrestle with difficult problems on their own. Schools should "
     "require students to complete at least one draft without any AI assistance of any "
     "kind, which would teach them the patience that genuine learning demands. Only then "
     "will they understand that struggle is not failure but the beginning of real growth."),
    ("Sample answer 3 (All-cleft) ",
     "All schools need to do is teach students how to use AI as a tool rather than a "
     "replacement for thinking. Many teenagers now trust an app more than their own "
     "judgment, which means they never develop the confidence to solve problems alone. A "
     "short practical lesson on verifying AI answers with real sources would show students "
     "that true intelligence is not about speed but about understanding. This small change "
     "could transform the way they approach every subject."),
]

SAMPLE_NOTE = ("Each sample uses ONE cleft: Sample 1 an it-cleft (It is \u2026 that \u2026), "
               "Sample 2 a wh-cleft (What \u2026 is \u2026) and Sample 3 an all-cleft (All "
               "\u2026 need to do is \u2026). Each sample also contains at least ONE relative "
               "clause or inverted pattern \u2014 point these out to students before they write.")


# ---------------------------------------------------------------- helpers
class Builder:
    def __init__(self):
        self.model = docx.Document(MODEL)
        self.M = list(self.model.element.body.iterchildren())
        self.doc = docx.Document(MODEL)
        body = self.doc.element.body
        self.sectPr = body.find(qn("w:sectPr"))
        for child in list(body.iterchildren()):
            if child is not self.sectPr:
                body.remove(child)
        self.body = body

    def clone(self, idx):
        return copy.deepcopy(self.M[idx])

    def push(self, el):
        self.body.insert(list(self.body).index(self.sectPr), el)
        return el

    @staticmethod
    def _rpr(fmt):
        rPr = etree.Element(qn("w:rPr"))
        rf = etree.SubElement(rPr, qn("w:rFonts"))
        rf.set(qn("w:cs"), "Times New Roman")
        if fmt in ("bold", "bold_u"):
            etree.SubElement(rPr, qn("w:b"))
            etree.SubElement(rPr, qn("w:bCs"))
        if fmt == "italic":
            etree.SubElement(rPr, qn("w:i"))
        if fmt == "bold_u":
            u = etree.SubElement(rPr, qn("w:u"))
            u.set(qn("w:val"), "single")
        return rPr

    def set_text(self, p, text):
        runs = p.findall(qn("w:r"))
        assert runs, "paragraph has no run to copy"
        for r in runs[1:]:
            p.remove(r)
        r = runs[0]
        for child in list(r):
            if child.tag != qn("w:rPr"):
                r.remove(child)
        t = etree.SubElement(r, qn("w:t"))
        t.text = text
        t.set(XMLSPACE, "preserve")

    def set_runs(self, p, parts):
        for r in p.findall(qn("w:r")):
            p.remove(r)
        for text, fmt in parts:
            r = etree.SubElement(p, qn("w:r"))
            r.append(self._rpr("plain" if fmt is None or fmt == "br" else fmt))
            if fmt == "br":
                etree.SubElement(r, qn("w:br"))
                continue
            if fmt is None:
                sym = etree.SubElement(r, qn("w:sym"))
                sym.set(qn("w:font"), "Wingdings")
                sym.set(qn("w:char"), "F0E0")
                continue
            t = etree.SubElement(r, qn("w:t"))
            t.text = text
            t.set(XMLSPACE, "preserve")
        return p

    def cell(self, tbl, r, c):
        tr = tbl.findall(qn("w:tr"))[r]
        tc = tr.findall(qn("w:tc"))[c]
        return tc.find(qn("w:p"))

    def set_cell(self, tbl, r, c, text):
        self.set_text(self.cell(tbl, r, c), text)

    def add_row(self, tbl, model_row_idx):
        tbl.append(copy.deepcopy(tbl.findall(qn("w:tr"))[model_row_idx]))


def build(variant="teacher"):
    if variant != "teacher":
        raise SystemExit("student copy is derived from the teacher docx")
    b = Builder()

    # -- header block
    for idx, text in ((H_SCHOOL, SCHOOL), (H_SUBJ, SUBJ), (H_TITLE, TITLE), (H_TOPIC, TOPIC)):
        el = b.clone(idx)
        b.set_text(el, text)
        b.push(el)

    # -- Part 1: Proofreading
    el = b.clone(P_H2); b.set_text(el, "Part 1: Proofreading quiz"); b.push(el)
    el = b.clone(P_PLAIN); b.set_text(el, QUIZ_INTRO); b.push(el)
    el = b.clone(P_PLAIN); b.set_text(el, PARA_A); b.push(el)
    el = b.clone(P_PLAIN); b.set_text(el, PARA_B); b.push(el)
    b.push(b.clone(P_BLANK))
    el = b.clone(P_ANS_LABEL); b.set_text(el, "Teacher\u2019s answers"); b.push(el)

    at = b.clone(T_ANSWERS)
    # Update header row
    b.set_cell(at, 0, 1, "Correction")
    while len(at.findall(qn("w:tr"))) < len(QUIZ_KEY) + 1:
        b.add_row(at, 1)
    for i, (num, corr) in enumerate(QUIZ_KEY, start=1):
        b.set_cell(at, i, 0, num)
        b.set_cell(at, i, 1, corr)
    b.push(at)
    b.push(b.clone(P_BLANK))

    # -- Part 2: Cleft patterns + class practice
    el = b.clone(P_H2B); b.set_text(el, "Part 2: Cleft sentences in formal writing"); b.push(el)
    el = b.clone(P_PLAIN2); b.set_text(el, P2_INTRO); b.push(el)
    for (title, plain, formal, demo), class_questions in zip(EXAMPLES, CLASS_PRACTICE):
        el = b.clone(P_H3); b.set_text(el, title); b.push(el)
        t = b.clone(T_PAIR)
        b.set_cell(t, 0, 0, "Plain:"); b.set_cell(t, 0, 1, plain)
        b.set_cell(t, 1, 0, "More formal:"); b.set_cell(t, 1, 1, formal)
        b.push(t)
        el = b.clone(P_BOX); b.set_text(el, demo); b.push(el)
        # Class practice questions (2 per pattern)
        for cp_title, simple, prompt, answer in class_questions:
            el = b.clone(P_SETTITLE); b.set_text(el, cp_title); b.push(el)
            t = b.clone(T_SET)
            b.set_cell(t, 0, 1, simple)
            b.set_cell(t, 1, 1, prompt)
            if variant == "teacher":
                b.set_cell(t, 2, 1, answer)
            else:
                b.set_cell(t, 2, 0, "")
                b.set_cell(t, 2, 1, "")
            b.push(t)

    # -- Part 3: Independent practice sets (4 only, no hints)
    el = b.clone(P_H1); b.set_text(el, "Part 3: Practice sets"); b.push(el)
    el = b.clone(P_SETINTRO); b.set_text(el, SETS_INTRO); b.push(el)
    for title, simple, prompt, answer in SETS:
        el = b.clone(P_SETTITLE); b.set_text(el, title); b.push(el)
        t = b.clone(T_SET)
        b.set_cell(t, 0, 1, simple)
        b.set_cell(t, 1, 1, prompt)
        if variant == "teacher":
            b.set_cell(t, 2, 1, answer)
        else:
            b.set_cell(t, 2, 0, "")
            b.set_cell(t, 2, 1, "")
        b.push(t)

    # -- Part 4: Writing Task
    el = b.clone(P_H1W); b.set_text(el, TASK_TITLE); b.push(el)
    el = b.clone(P_INSTR_B); b.set_text(el, "Instructions:"); b.push(el)
    el = b.clone(P_INSTR_1); b.set_text(el, TASK_LINE); b.push(el)
    el = b.clone(P_INSTR_2); b.set_runs(el, TASK_BULLETS); b.push(el)
    b.push(b.clone(P_BLANK))
    for label, body in SAMPLES:
        el = b.clone(P_SAMPLE)
        b.set_runs(el, [(label, "bold"), (body, "plain")])
        b.push(el)
    el = b.clone(P_NOTE); b.set_text(el, SAMPLE_NOTE); b.push(el)

    out = OUT[variant]
    b.doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "teacher")
